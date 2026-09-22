"""File guide.

- Use: Tests the shared deterministic document extraction boundary.
- Usage: Run with pytest when changing extraction, upload ingestion, or attachments.
- Duties: Builds small fixtures for each public upload format and verifies
  normalized RAG documents with source-location metadata and safe errors.
- Depends on: External packages docx, openpyxl, and project RAG extraction.
"""

import json
from pathlib import Path

import pytest
from docx import Document as DOCXDocument
from openpyxl import Workbook

from openagentlab.rag.extraction import (
    EMPTY_DOCUMENT,
    MALFORMED_DOCX,
    MALFORMED_JSON,
    MALFORMED_XLSX,
    DefaultDocumentExtractor,
    DocumentExtractionError,
    SharedDocumentLoader,
)


def test_txt_and_markdown_extract_line_ranges_from_bytes() -> None:
    extractor = DefaultDocumentExtractor()

    txt = extractor.extract_bytes(
        b"alpha\nbeta\n",
        filename="notes.txt",
        source_id="doc-1",
    )
    markdown = extractor.extract_bytes(
        b"# Heading\nbody\n",
        filename="notes.md",
        source_id="doc-2",
    )

    assert txt.documents[0].text == "alpha\nbeta"
    assert txt.documents[0].metadata["line_start"] == 1
    assert txt.documents[0].metadata["line_end"] == 2
    assert txt.documents[0].metadata["document_id"] == "doc-1"
    assert markdown.documents[0].metadata["file_type"] == "md"
    assert markdown.documents[0].metadata["source_location"] == "lines:1-2"


def test_pdf_extracts_page_metadata(tmp_path: Path) -> None:
    path = tmp_path / "fixture.pdf"
    write_text_pdf(path, ("first page", "second page"))

    result = DefaultDocumentExtractor().extract_path(path, source_id="doc-pdf")

    assert [document.metadata["page_number"] for document in result.documents] == [
        1,
        2,
    ]
    assert result.documents[0].metadata["page_count"] == 2
    assert result.documents[0].metadata["source_location"] == "page:1"
    assert result.documents[0].metadata["document_id"] == "doc-pdf"


def test_csv_extracts_row_ranges_and_columns_from_bytes() -> None:
    result = DefaultDocumentExtractor().extract_bytes(
        b"name,age\nAlice,32\nBob,41\n",
        filename="people.csv",
        source_id="doc-csv",
    )

    document = result.documents[0]
    assert document.text.startswith("Columns: name, age")
    assert "Row 2: name=Alice | age=32" in document.text
    assert document.metadata["location_type"] == "row_range"
    assert document.metadata["row_start"] == 2
    assert document.metadata["row_end"] == 3
    assert document.metadata["columns"] == ["name", "age"]


def test_json_extracts_json_path_segments_from_bytes() -> None:
    content = json.dumps(
        {"items": [{"name": "alpha"}, {"name": "beta"}], "count": 2}
    ).encode()

    result = DefaultDocumentExtractor().extract_bytes(
        content,
        filename="data.json",
        source_id="doc-json",
    )

    paths = {document.metadata["json_path"] for document in result.documents}
    assert "$.items" in paths
    assert "$.count" in paths
    assert all(
        document.metadata["location_type"] == "json_path"
        for document in result.documents
    )


def test_docx_extracts_paragraphs_and_table_rows_from_bytes(tmp_path: Path) -> None:
    content = _docx_bytes(tmp_path)

    result = DefaultDocumentExtractor().extract_bytes(
        content,
        filename="report.docx",
        source_id="doc-docx",
    )

    paragraph = next(
        document
        for document in result.documents
        if document.metadata["location_type"] == "paragraph"
    )
    table = next(
        document
        for document in result.documents
        if document.metadata["location_type"] == "table_rows"
    )
    assert paragraph.metadata["paragraph_index"] == 1
    assert paragraph.metadata["source_location"] == "paragraph:1"
    assert table.metadata["table_index"] == 1
    assert table.metadata["row_start"] == 2
    assert table.metadata["columns"] == ["name", "value"]


def test_xlsx_extracts_sheet_row_ranges_and_skips_empty_sheets(tmp_path: Path) -> None:
    content = _xlsx_bytes(tmp_path)

    result = DefaultDocumentExtractor().extract_bytes(
        content,
        filename="workbook.xlsx",
        source_id="doc-xlsx",
    )

    sheet_names = {document.metadata["sheet_name"] for document in result.documents}
    assert sheet_names == {"Summary", "Details"}
    summary = result.documents[0]
    assert summary.metadata["location_type"] == "sheet_rows"
    assert summary.metadata["row_start"] == 2
    assert summary.metadata["row_end"] == 2
    assert summary.metadata["columns"] == ["name", "value"]
    assert summary.metadata["source_location"] == "sheet:Summary:rows:2-2"


def test_shared_loader_uses_same_extractor_contract(tmp_path: Path) -> None:
    path = tmp_path / "data.json"
    path.write_text('{"answer": 42}', encoding="utf-8")

    documents = SharedDocumentLoader(source_id="doc-loader").load(path)

    assert documents[0].metadata["document_id"] == "doc-loader"
    assert documents[0].metadata["json_path"] == "$.answer"


def test_malformed_json_becomes_safe_extraction_error() -> None:
    with pytest.raises(DocumentExtractionError) as exc_info:
        DefaultDocumentExtractor().extract_bytes(b'{"broken": ', filename="bad.json")

    assert exc_info.value.code == MALFORMED_JSON
    assert "Malformed JSON file" in exc_info.value.safe_message


def test_corrupt_docx_and_xlsx_become_safe_extraction_errors() -> None:
    extractor = DefaultDocumentExtractor()

    with pytest.raises(DocumentExtractionError) as docx_error:
        extractor.extract_bytes(b"not a docx", filename="bad.docx")
    with pytest.raises(DocumentExtractionError) as xlsx_error:
        extractor.extract_bytes(b"not a workbook", filename="bad.xlsx")

    assert docx_error.value.code == MALFORMED_DOCX
    assert xlsx_error.value.code == MALFORMED_XLSX


def test_pdf_with_no_extractable_text_fails_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    write_text_pdf(path, ("",))

    with pytest.raises(DocumentExtractionError) as exc_info:
        DefaultDocumentExtractor().extract_path(path)

    assert exc_info.value.code == EMPTY_DOCUMENT


def _docx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "report.docx"
    document = DOCXDocument()
    document.add_paragraph("Executive summary")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "42"
    document.save(path)
    return path.read_bytes()


def _xlsx_bytes(tmp_path: Path) -> bytes:
    path = tmp_path / "workbook.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["name", "value"])
    summary.append(["alpha", 1])
    details = workbook.create_sheet("Details")
    details.append(["id", "status"])
    details.append(["A-1", "done"])
    workbook.create_sheet("Empty")
    workbook.save(path)
    workbook.close()
    return path.read_bytes()


def write_text_pdf(path: Path, page_texts: tuple[str, ...]) -> None:
    objects: list[bytes] = []
    page_numbers: list[int] = []

    def add_object(content: bytes) -> int:
        objects.append(content)
        return len(objects)

    add_object(b"<< /Type /Catalog /Pages 2 0 R >>")
    add_object(b"")
    add_object(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    for text in page_texts:
        page_number = len(objects) + 1
        content_number = len(objects) + 2
        page_numbers.append(page_number)

        add_object(
            (
                f"<< /Type /Page /Parent 2 0 R "
                f"/MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 3 0 R >> >> "
                f"/Contents {content_number} 0 R >>"
            ).encode()
        )

        content = f"BT\n/F1 12 Tf\n72 720 Td\n{_pdf_string(text)} Tj\nET\n".encode()
        add_object(
            b"<< /Length "
            + str(len(content)).encode()
            + b" >>\nstream\n"
            + content
            + b"endstream"
        )

    kids = " ".join(f"{page_number} 0 R" for page_number in page_numbers)
    objects[1] = (
        f"<< /Type /Pages /Kids [{kids}] /Count {len(page_numbers)} >>".encode()
    )
    _write_pdf_objects(path, objects)


def _pdf_string(value: str) -> str:
    escaped = value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return f"({escaped})"


def _write_pdf_objects(path: Path, objects: list[bytes]) -> None:
    content = b"%PDF-1.4\n"
    offsets = []

    for object_number, object_content in enumerate(objects, start=1):
        offsets.append(len(content))
        content += f"{object_number} 0 obj\n".encode() + object_content + b"\nendobj\n"

    xref_start = len(content)
    content += f"xref\n0 {len(objects) + 1}\n".encode()
    content += b"0000000000 65535 f \n"

    for offset in offsets:
        content += f"{offset:010d} 00000 n \n".encode()

    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R"
    trailer += f" >>\nstartxref\n{xref_start}\n%%EOF\n"

    path.write_bytes(content + trailer.encode())
